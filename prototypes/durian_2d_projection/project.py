#!/usr/bin/env python3
"""PROTOTYPE: project turntable videos onto disposable spherical UV atlases."""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from scipy import ndimage


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).resolve().parent / "output"
FRAME_WIDTH = 960
FRAME_HEIGHT = 540
SAMPLES = 72
ATLAS_WIDTH = 1440
ATLAS_HEIGHT = 512
SURFACE_STEP = 2
MAX_VIEW_ANGLE = np.deg2rad(24)
MIN_TEXTURE_MOTION_DEGREES = 60


@dataclass(frozen=True)
class Video:
    fruit: str
    orientation: str
    view: str
    archive: Path
    member: str
    slug: str


def normalized(value: str) -> str:
    return unicodedata.normalize("NFC", value)


def fold(value: str) -> str:
    value = normalized(value).replace("đ", "d").replace("Đ", "D")
    return "".join(
        char
        for char in unicodedata.normalize("NFD", value)
        if unicodedata.category(char) != "Mn"
    ).lower()


def sample_key(value: str) -> tuple[str, int, int]:
    match = re.fullmatch(r"([A-Z]+)(\d+)(?:C(\d+))?", value.upper())
    return (match.group(1), int(match.group(2)), int(match.group(3) or 0)) if match else (value, 0, 0)


def discover(root: Path) -> list[Video]:
    records: list[tuple[str, str, str, Path, str]] = []
    archives = [root] if root.is_file() else [
        path
        for path in root.rglob("*.zip")
        if not any(part.startswith((".", "$")) for part in path.relative_to(root).parts)
    ]
    for archive in sorted(archives):
        name = archive.name.lower()
        if "iphone" in name or ("lidar" in name and "rename" not in name):
            continue
        with zipfile.ZipFile(archive) as bundle:
            for info in bundle.infolist():
                member = normalized(info.filename)
                if Path(member).name.startswith("._"):
                    continue
                fruit_match = re.search(r"\b(Q\d+|V\d+C\d+)\b", member, re.IGNORECASE)
                if not fruit_match:
                    continue
                fruit = fruit_match.group(1).upper()
                plain = fold(member)
                if member.lower().endswith(".mov"):
                    orientation = "ngang" if "/ngang/" in plain else "dung"
                    if "(tren" in plain:
                        view = "tren"
                    elif "(duoi" in plain:
                        view = "duoi"
                    else:
                        view = "unknown"
                elif member.lower().endswith("/rgb.mp4"):
                    orientation = "ngang" if " ngang" in plain else "dung"
                    date = "21" if "21jul" in name else "18"
                    background = "-trang" if "trang" in plain else "-den" if "den" in plain else ""
                    view = f"lidar-{date}{background}"
                else:
                    continue
                records.append((fruit, orientation, view, archive, info.filename))

    workbook = root / "RENAME.xlsx" if root.is_dir() else None
    if workbook and workbook.is_file():
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError("RENAME.xlsx requires openpyxl") from error

        sheets = {
            "IP11.20.7.Whitefont": ("Video-iphone11-20Jul2026.zip", "iphone11-20-trang"),
            "IP11.20.7.Blackfont": ("Video-iphone11-20Jul2026.zip", "iphone11-20-den"),
            "IP11.18.7": ("Video-iphone11-18Jul2026.zip", "iphone11-18"),
            "IP12.20.7.Blackfont": ("Video-iphone12-20Jul2026.zip", "iphone12-20-den"),
            "IP12.20.7.whitefont": ("Video-iphone12-20Jul2026.zip", "iphone12-20-trang"),
            "IP12.18.7": ("Video-iphone12-18Jul2026.zip", "iphone12-18"),
        }
        mapped: dict[tuple[Path, str], set[tuple[str, str, str]]] = {}
        book = load_workbook(workbook, read_only=True, data_only=True)
        for sheet in book.worksheets:
            config = sheets.get(sheet.title.strip())
            if not config:
                continue
            archive, view = root / config[0], config[1]
            if not archive.is_file():
                continue
            with zipfile.ZipFile(archive) as bundle:
                members = {
                    Path(info.filename).name.upper(): info.filename
                    for info in bundle.infolist()
                    if not Path(info.filename).name.startswith("._")
                }
            for row in sheet.iter_rows(min_row=2, values_only=True):
                image, sample, direction = row[1:3] + row[5:6]
                match = re.fullmatch(r"V\d+C\d+", str(sample).upper()) if sample else None
                if not image or not match or fold(str(direction)) not in {"dung", "ngang"}:
                    continue
                member = members.get(f"IMG_{int(image):04d}.MOV")
                if member:
                    mapped.setdefault((archive, member), set()).add(
                        (match.group().upper(), fold(str(direction)), view)
                    )
        for (archive, member), labels in mapped.items():
            if len(labels) == 1:  # Ambiguous spreadsheet rows are unsafe to guess.
                fruit, orientation, view = labels.pop()
                records.append((fruit, orientation, view, archive, member))

    for fruit in {record[0] for record in records}:
        if not fruit.startswith("Q"):
            continue
        for orientation in ("ngang", "dung"):
            indices = [
                index
                for index, record in enumerate(records)
                if record[0] == fruit and record[1] == orientation
            ]
            views = [records[index][2] for index in indices]
            if len(indices) != 2 or set(views) == {"tren", "duoi"}:
                continue
            missing = "duoi" if "tren" in views else "tren"
            replace = next(
                (index for index in reversed(indices) if records[index][2] == "unknown"),
                indices[-1],
            )
            old = records[replace]
            records[replace] = (old[0], old[1], missing, old[3], old[4])

    counts: dict[str, int] = {}
    videos: list[Video] = []
    for fruit, orientation, view, archive, member in sorted(
        records,
        key=lambda item: (
            sample_key(item[0]),
            item[1] != "ngang",
            item[2],
            item[3].name,
            item[4],
        ),
    ):
        base = f"{fruit}-{orientation}-{view}"
        counts[base] = counts.get(base, 0) + 1
        suffix = f"-{counts[base]}" if counts[base] > 1 else ""
        videos.append(Video(fruit, orientation, view, archive, member, base + suffix))
    return videos


def run(command: list[str], *, stdin: bytes | None = None) -> bytes:
    result = subprocess.run(
        command,
        input=stdin,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stderr.decode("utf-8", errors="replace"))
    return result.stdout


def probe(path: Path) -> dict:
    raw = run(
        [
            "ffprobe",
            "-v",
            "error",
            "-select_streams",
            "v:0",
            "-show_entries",
            "stream=codec_name,width,height,avg_frame_rate:format=duration,size",
            "-of",
            "json",
            str(path),
        ]
    )
    return json.loads(raw)


def decode_frames(
    path: Path,
    duration: float,
    *,
    width: int = FRAME_WIDTH,
    height: int = FRAME_HEIGHT,
    samples: int = SAMPLES,
) -> np.ndarray:
    start = duration * 0.02
    usable = duration * 0.96
    rate = samples / usable
    command = ["ffmpeg", "-v", "error"]
    if sys.platform == "darwin":
        command += ["-hwaccel", "videotoolbox"]
    raw = run(
        command
        + [
            "-ss",
            f"{start:.6f}",
            "-i",
            str(path),
            "-t",
            f"{usable:.6f}",
            "-vf",
            (
                f"fps={rate:.9f},"
                f"scale={width}:{height}:force_original_aspect_ratio=decrease,"
                f"pad={width}:{height}:(ow-iw)/2:(oh-ih)/2"
            ),
            "-an",
            "-f",
            "rawvideo",
            "-pix_fmt",
            "rgb24",
            "pipe:1",
        ]
    )
    pixels_per_frame = width * height * 3
    count = len(raw) // pixels_per_frame
    if not count:
        raise RuntimeError("ffmpeg returned no frames")
    frames = np.frombuffer(raw[: count * pixels_per_frame], dtype=np.uint8)
    frames = frames.reshape(count, height, width, 3)
    if count > samples:
        frames = frames[:samples]
    return frames


def active_frames(frames: np.ndarray) -> np.ndarray:
    height, width = frames.shape[1:3]
    sample = frames[
        :, int(height * 0.11) : int(height * 0.93) : 4,
        int(width * 0.10) : int(width * 0.90) : 4,
    ].astype(np.int16)
    motion = np.abs(sample[1:] - sample[:-1]).mean(axis=(1, 2, 3))
    threshold = max(2.5, float(np.percentile(motion, 75)) * 0.35)
    moving = np.flatnonzero(motion > threshold)
    if not len(moving):
        return frames
    start = max(0, int(moving[0]))
    end = min(len(frames), int(moving[-1]) + 2)
    return frames[start:end]


def fruit_mask(frame: np.ndarray) -> np.ndarray | None:
    height, width = frame.shape[:2]
    rgb = frame.astype(np.int16)
    red, green, blue = rgb[..., 0], rgb[..., 1], rgb[..., 2]
    chroma = rgb.max(axis=2) - rgb.min(axis=2)
    mask = (
        (green - blue > 7)
        & (red - blue > 2)
        & (chroma > 11)
        & (rgb.mean(axis=2) > 55)
    )
    mask[: int(height * 0.06)] = False
    mask[:, : int(width * 0.05)] = False
    mask[:, int(width * 0.95) :] = False
    for row in range(height):
        xs = np.flatnonzero(mask[row])
        if len(xs) and xs[-1] - xs[0] > width * 0.78:
            mask[row] = False
    mask = ndimage.binary_opening(mask, iterations=1)
    mask = ndimage.binary_closing(mask, iterations=3)

    labels, count = ndimage.label(mask)
    if not count:
        return None

    best_label = 0
    best_score = 0.0
    image_center = np.array([height * 0.5, width * 0.5])
    for label_id in range(1, count + 1):
        ys, xs = np.where(labels == label_id)
        area = len(xs)
        if area < 2_000:
            continue
        center = np.array([ys.mean(), xs.mean()])
        distance = np.linalg.norm((center - image_center) / image_center)
        touches_edge = (
            xs.min() <= 1
            or xs.max() >= width - 2
            or ys.min() <= 1
            or ys.max() >= height - 2
        )
        score = area / (1 + distance * 2)
        if touches_edge:
            score *= 0.15
        if score > best_score:
            best_score = score
            best_label = label_id

    if not best_label:
        return None
    result = labels == best_label
    result = ndimage.binary_fill_holes(result)
    result = ndimage.binary_closing(result, iterations=4)
    result = ndimage.binary_dilation(result, iterations=1)

    _, top, _, bottom = bounds(result)
    # ponytail: discard the contact zone; detect the table line if cap coverage matters.
    result[top + int((bottom - top) * 0.88) :] = False
    return result


def bounds(mask: np.ndarray) -> tuple[int, int, int, int]:
    ys, xs = np.where(mask)
    return int(xs.min()), int(ys.min()), int(xs.max()) + 1, int(ys.max()) + 1


def frame_phases(valid: list[tuple[np.ndarray, np.ndarray]]) -> np.ndarray:
    """Measure turntable angle from horizontal texture motion."""
    size = 192
    textures = []
    for frame, mask in valid:
        x0, y0, x1, y1 = bounds(mask)
        gray = frame[y0:y1, x0:x1].mean(axis=2)
        gray = ndimage.zoom(
            gray, (size / gray.shape[0], size / gray.shape[1]), order=1
        )[:size, :size]
        gray = ndimage.sobel(gray, axis=1)
        gray[:15] = gray[-15:] = 0
        gray[:, :20] = gray[:, -20:] = 0
        gray = (gray - gray.mean()) / (gray.std() + 1e-6)
        gray *= np.outer(np.hanning(size), np.hanning(size))
        textures.append(gray)

    shifts = []
    for before, after in zip(textures, textures[1:]):
        cross = np.fft.fft2(before) * np.conj(np.fft.fft2(after))
        cross /= np.maximum(np.abs(cross), 1e-9)
        correlation = np.fft.ifft2(cross).real
        _, shift_x = np.unravel_index(np.argmax(correlation), correlation.shape)
        if shift_x > size // 2:
            shift_x -= size
        shifts.append(float(shift_x))

    if not shifts:
        return np.zeros(len(valid))
    shifts = np.asarray(shifts)
    moving = shifts[np.abs(shifts) > 1]
    direction = np.sign(np.median(moving)) if len(moving) else -1
    increments = shifts / (size / 2)
    typical = np.median(np.abs(increments[increments != 0])) if np.any(increments) else 0
    increments[
        (np.sign(increments) != direction)
        | (np.abs(increments) > max(typical * 2.5, 0.2))
    ] = 0
    phases = np.concatenate(([0.0], np.cumsum(increments)))
    return phases


def spherical_atlas(
    valid: list[tuple[np.ndarray, np.ndarray]], phases: np.ndarray
) -> np.ndarray:
    """Inverse-project overlapping camera pixels onto longitude/latitude UV."""
    pixels = ATLAS_HEIGHT * ATLAS_WIDTH
    best_weight = np.full(pixels, -1.0, dtype=np.float32)
    atlas = np.zeros((pixels, 3), dtype=np.uint8)

    for phase, (frame, mask) in zip(phases, valid):
        x0, y0, x1, y1 = bounds(mask)
        ys = np.arange(y0, y1, SURFACE_STEP)
        xs = np.arange(x0, x1, SURFACE_STEP)
        if not len(ys) or not len(xs):
            continue

        sampled_mask = mask[np.ix_(ys, xs)]
        left = np.zeros(len(ys), dtype=np.float32)
        right = np.zeros(len(ys), dtype=np.float32)
        row_ok = np.zeros(len(ys), dtype=bool)
        for row_index, y in enumerate(ys):
            occupied = np.flatnonzero(mask[y, x0:x1]) + x0
            if len(occupied) >= 3:
                left[row_index], right[row_index] = occupied[0], occupied[-1]
                row_ok[row_index] = True

        center_x = (left + right)[:, None] / 2
        radius_x = np.maximum((right - left)[:, None] / 2, 1)
        normalized_x = (xs[None, :] - center_x) / radius_x
        longitude_offset = np.arcsin(np.clip(normalized_x, -1, 1))
        usable = sampled_mask & row_ok[:, None] & (np.abs(longitude_offset) <= MAX_VIEW_ANGLE)
        if not usable.any():
            continue

        # An orthographic camera sees y = sin(latitude) on an ellipsoid.
        center_y = (y0 + y1 - 1) / 2
        radius_y = max((y1 - y0 - 1) / 2, 1)
        sin_latitude = np.clip((center_y - ys) / radius_y, -1, 1)
        # Cylindrical equal-area: unlike equirectangular/Mercator, it does not
        # explode the sparse measurements around the fruit's two poles.
        atlas_y = np.rint((0.5 - sin_latitude / 2) * (ATLAS_HEIGHT - 1)).astype(int)
        atlas_x = np.floor(
            np.mod(phase + longitude_offset, 2 * np.pi) / (2 * np.pi) * ATLAS_WIDTH
        ).astype(int)
        flat_index = atlas_y[:, None] * ATLAS_WIDTH + atlas_x

        selected = usable.ravel()
        indexes = flat_index.ravel()[selected]
        # Favor the least-oblique observation to avoid blurred spike silhouettes.
        confidence = np.cos(longitude_offset).ravel()[selected] ** 64
        source = frame[np.ix_(ys, xs)].reshape(-1, 3)[selected]
        better = confidence > best_weight[indexes]
        indexes = indexes[better]
        best_weight[indexes] = confidence[better]
        atlas[indexes] = source[better]

    covered = best_weight >= 0
    if not covered.any():
        raise RuntimeError("spherical projection produced no pixels")
    atlas = atlas.reshape(ATLAS_HEIGHT, ATLAS_WIDTH, 3)

    holes = ~covered.reshape(ATLAS_HEIGHT, ATLAS_WIDTH)
    if holes.any():
        # ponytail: nearest fill is enough for inspection; use seam-aware inpainting
        # only if this projection passes repeat-capture matching.
        distance, nearest = ndimage.distance_transform_edt(
            holes, return_distances=True, return_indices=True
        )
        small_holes = holes & (distance <= 4)
        atlas[small_holes] = atlas[nearest[0][small_holes], nearest[1][small_holes]]
    return atlas


def edge_map(atlas: np.ndarray) -> np.ndarray:
    gray = atlas.astype(np.float32).mean(axis=2)
    dy, dx = np.gradient(gray)
    magnitude = np.hypot(dx, dy)
    nonzero = magnitude[magnitude > 0]
    ceiling = np.percentile(nonzero, 96) if len(nonzero) else 1
    edges = np.clip(magnitude * 255 / max(ceiling, 1), 0, 255).astype(np.uint8)
    return np.repeat(edges[..., None], 3, axis=2)


def overlay(frame: np.ndarray, mask: np.ndarray | None) -> np.ndarray:
    result = frame.copy()
    if mask is None:
        return result
    boundary = mask ^ ndimage.binary_erosion(mask)
    result[boundary] = (255, 40, 40)
    return result


def write_png(image: np.ndarray, path: Path) -> None:
    height, width = image.shape[:2]
    pixel_format = "rgba" if image.shape[2] == 4 else "rgb24"
    run(
        [
            "ffmpeg",
            "-v",
            "error",
            "-f",
            "rawvideo",
            "-pix_fmt",
            pixel_format,
            "-s",
            f"{width}x{height}",
            "-i",
            "pipe:0",
            "-frames:v",
            "1",
            "-y",
            str(path),
        ],
        stdin=np.ascontiguousarray(image).tobytes(),
    )


def process(video: Video, output: Path, min_motion: float = MIN_TEXTURE_MOTION_DEGREES) -> dict:
    video_output = output / video.fruit
    video_output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="durian-2d-") as temporary:
        source = Path(temporary) / "source.mov"
        with zipfile.ZipFile(video.archive) as bundle:
            with bundle.open(video.member) as incoming, source.open("wb") as outgoing:
                shutil.copyfileobj(incoming, outgoing, length=8 * 1024 * 1024)

        metadata = probe(source)
        duration = float(metadata["format"]["duration"])
        decoded = decode_frames(source, duration)

    frames = active_frames(decoded)
    masks = [fruit_mask(frame) for frame in frames]
    valid = [(frame, mask) for frame, mask in zip(frames, masks) if mask is not None]
    if len(valid) < SAMPLES // 2:
        raise RuntimeError(f"fruit mask failed on {len(frames) - len(valid)}/{len(frames)} frames")

    measured_phases = frame_phases(valid)
    motion = float(abs(np.rad2deg(measured_phases[-1])))
    if motion < min_motion:
        raise RuntimeError(f"only {motion:.1f}° texture motion; need {min_motion:.0f}°")
    # ponytail: captures are one constant-speed turn; use encoder phases if that protocol changes.
    direction = np.sign(measured_phases[-1])
    phases = np.linspace(0, direction * 2 * np.pi, len(valid), endpoint=False)
    atlas = spherical_atlas(valid, phases)
    atlas_path = video_output / f"{video.slug}-atlas.png"
    edges_path = video_output / f"{video.slug}-edges.png"
    write_png(atlas, atlas_path)
    write_png(edge_map(atlas), edges_path)

    preview_indices = np.linspace(0, len(frames) - 1, 6).round().astype(int)
    previews = [overlay(frames[index], masks[index])[::2, ::2] for index in preview_indices]
    preview_path = video_output / f"{video.slug}-mask.png"
    write_png(np.concatenate(previews, axis=1), preview_path)

    boxes = [bounds(mask) for mask in masks if mask is not None]
    median_width = int(np.median([box[2] - box[0] for box in boxes]))
    median_height = int(np.median([box[3] - box[1] for box in boxes]))
    return {
        "fruit": video.fruit,
        "orientation": video.orientation,
        "view": video.view,
        "slug": video.slug,
        "archive": video.archive.name,
        "member": normalized(video.member),
        "duration_seconds": round(duration, 3),
        "decoded_frames": len(decoded),
        "active_frames": len(frames),
        "valid_masks": len(valid),
        "texture_motion_degrees": round(motion, 1),
        "median_fruit_width": median_width,
        "median_fruit_height": median_height,
        "atlas": atlas_path.relative_to(output).as_posix(),
        "edges": edges_path.relative_to(output).as_posix(),
        "mask_preview": preview_path.relative_to(output).as_posix(),
    }


def write_report(output: Path, rows: list[dict]) -> None:
    rows.sort(key=lambda row: (sample_key(row["fruit"]), row["orientation"], row["view"], row["slug"]))
    fields = list(rows[0]) if rows else []
    with (output / "manifest.csv").open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    sections = []
    fruits = sorted({row["fruit"] for row in rows}, key=sample_key)
    for fruit in fruits:
        cards = []
        for row in [item for item in rows if item["fruit"] == fruit]:
            label = html.escape(f"{row['orientation']} / {row['view']}")
            cards.append(
                f"""
                <article>
                  <h3>{label}</h3>
                  <p>{row['valid_masks']}/{row['active_frames']} active masks
                    ({row['decoded_frames']} sampled) · {row['duration_seconds']}s ·
                    {row['texture_motion_degrees']}° texture motion ·
                    ROI {row['median_fruit_width']}×{row['median_fruit_height']}</p>
                  <img src="{html.escape(row['atlas'])}" alt="Spherical UV atlas">
                  <img src="{html.escape(row['edges'])}" alt="edge fingerprint">
                  <details><summary>Mask check</summary>
                    <img class="wide" src="{html.escape(row['mask_preview'])}" alt="mask preview">
                  </details>
                </article>
                """
            )
        sections.append(f"<section><h2>{fruit}</h2><div class='grid'>{''.join(cards)}</div></section>")

    document = f"""<!doctype html>
<html lang="vi"><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>Durian 2D Projection Prototype</title>
<style>
body{{font:14px system-ui;margin:24px;background:#111;color:#eee}}h1,h2{{margin-top:32px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(360px,1fr));gap:18px}}
article{{background:#1d1d1d;padding:14px;border-radius:10px}}img{{width:100%;background:#000;margin-top:8px}}
.wide{{width:100%}}p,summary{{color:#aaa}}code{{background:#222;padding:2px 5px}}
</style>
<h1>PROTOTYPE — Durian 2D Projection</h1>
<p>{len(rows)} videos · {len(fruits)} fruits. Atlas: spherical longitude/latitude UV. Edge image: candidate fingerprint.</p>
{''.join(sections)}
</html>"""
    (output / "index.html").write_text(document, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("fruits", nargs="*", help="Optional fruit IDs, for example Q11 Q16")
    parser.add_argument("--source", type=Path, default=ROOT, help="Directory containing ZIP archives")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--jobs", type=int, default=2)
    parser.add_argument("--capture", action="append", help="Exact capture label; repeat to select several")
    parser.add_argument("--min-motion", type=float, default=MIN_TEXTURE_MOTION_DEGREES)
    parser.add_argument("--list", action="store_true", help="List discovered videos without processing")
    parser.add_argument("--self-check", action="store_true")
    args = parser.parse_args()

    if args.self_check:
        frame = np.zeros((120, 120, 3), dtype=np.uint8)
        yy, xx = np.ogrid[:120, :120]
        mask = ((xx - 60) / 42) ** 2 + ((yy - 60) / 50) ** 2 <= 1
        frame[mask] = (180, 140, 40)
        phases = np.linspace(0, 2 * np.pi, 12, endpoint=False)
        atlas = spherical_atlas([(frame, mask)] * 12, phases)
        assert atlas.shape == (ATLAS_HEIGHT, ATLAS_WIDTH, 3)
        assert np.count_nonzero(atlas) > atlas.size * 0.3
        assert sample_key("V12C3") < sample_key("V13C1")
        print("OK: spherical UV projection")
        return

    wanted = {fruit.upper() for fruit in args.fruits}
    videos = [video for video in discover(args.source) if not wanted or video.fruit in wanted]
    if args.capture:
        videos = [video for video in videos if video.view in args.capture]
    if args.list:
        for video in videos:
            print(f"{video.fruit}\t{video.orientation}\t{video.view}\t{video.archive.name}\t{video.member}")
        print(f"{len(videos)} videos · {len({video.fruit for video in videos})} samples")
        return

    missing = [tool for tool in ("ffmpeg", "ffprobe") if not shutil.which(tool)]
    if missing:
        raise SystemExit(f"Missing required tools: {', '.join(missing)}")

    if not videos:
        raise SystemExit("No matching videos found in ZIP archives")

    args.output.mkdir(parents=True, exist_ok=True)
    rows = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, args.jobs)) as executor:
        pending = {
            executor.submit(process, video, args.output, args.min_motion): video for video in videos
        }
        for index, future in enumerate(concurrent.futures.as_completed(pending), 1):
            video = pending[future]
            try:
                rows.append(future.result())
                print(f"[{index}/{len(videos)}] {video.slug}", flush=True)
            except Exception as exc:
                print(f"[{index}/{len(videos)}] {video.slug} skipped: {exc}", flush=True)

    write_report(args.output, rows)
    print(f"Report: {args.output / 'index.html'}")


if __name__ == "__main__":
    main()
