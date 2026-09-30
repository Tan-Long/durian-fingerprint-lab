"""Run: python3 -B -m unittest scripts.test_audit_morphology"""

import csv
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from scripts.audit_morphology import audit, validate_output, FRUIT_SHEET, LOCULE_SHEET, REQUIRED


class AuditMorphologyTests(unittest.TestCase):
    def fixture(self, directory, fruits=None, locules=None):
        path = Path(directory) / "labels.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        rows = {
            FRUIT_SHEET: fruits if fruits is not None else [["F1", "N1", "V1", "C1", 2, 1, 12, 0]],
            LOCULE_SHEET: locules if locules is not None else [["F1", 1, 1], ["F1", 2, 0]],
        }
        for name, values in rows.items():
            sheet = workbook.create_sheet(name)
            sheet.append(REQUIRED[name])
            sheet.append(["mm" if field == "dày vỏ" else None for field in REQUIRED[name]])
            for row in values:
                sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_valid_variable_count_zero_and_source_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            for count in [2, 3, 4, 5, 6]:
                path = self.fixture(directory, fruits=[["F1", "N1", "V1", "C1", count, count - 1, 12, 0]],
                                    locules=[["F1", n, int(n != count)] for n in range(1, count + 1)])
                before = path.read_bytes()
                report = audit(path)
                self.assertEqual(report["issues"], [])
                self.assertEqual(report["summary"]["valid_targets"]["locule_count"], 1)
                self.assertEqual(report["fruits"][0]["targets"]["empty_locule_count"]["value"], 0)
                self.assertEqual(path.read_bytes(), before)

    def test_missing_does_not_become_zero_or_invalidate_shell(self):
        with tempfile.TemporaryDirectory() as directory:
            report = audit(self.fixture(directory, fruits=[["F1", "N1", "V1", "C1", 2, None, 12, None]]))
            targets = report["fruits"][0]["targets"]
            self.assertIsNone(targets["aril_count"]["value"])
            self.assertIn("missing_numeric", targets["empty_locule_count"]["reasons"])
            self.assertEqual(targets["shell_thickness_mm"]["status"], "valid")

    def test_malformed_and_duplicate_keys_have_cell_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            report = audit(self.fixture(directory, locules=[
                ["F1", "F1", 2], ["F1", 0, 1], ["F1", 1.5, 1],
                ["F1", 1, 1], ["F1", 1, 1], [None, 2, 0], ["unknown", 1, None],
            ]))
            codes = report["summary"]["issues_by_code"]
            for code in ["nonnumeric", "nonpositive_numeric", "noninteger_numeric", "duplicate_locule_key",
                         "missing_or_invalid_fruit_id", "unknown_detail_fruit_id", "locule_count_mismatch", "aril_sum_mismatch"]:
                self.assertIn(code, codes)
            malformed = next(i for i in report["issues"] if i["code"] == "nonnumeric")
            self.assertEqual(malformed["sources"], [{"sheet": LOCULE_SHEET, "row": 3, "cells": ["B3"]}])
            self.assertEqual(report["fruits"][0]["targets"]["shell_thickness_mm"]["status"], "valid")
            self.assertEqual(report["summary"]["locule_rows_without_valid_fruit_id"], 1)
            self.assertEqual(report["summary"]["valid_locule_number_rows_with_fruit_id"], 3)

    def test_exact_crosswalk_no_bonus_repair_and_duplicate_identity(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.fixture(directory, fruits=[
                ["F1", "N1", "V1", "C1", 2, 1, 12, 0],
                ["bonus", "N1", "V4", "C1 bonus", 2, 1, 12, None],
                ["duplicate", "N2", "V1", "C1", 2, 1, 12, None],
                ["duplicate", "N2", "V1", "C1", 2, 1, 12, None],
            ])
            samples = Path(directory) / "samples.csv"
            with samples.open("w", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["collection_id", "sample_id", "sample_uid"])
                for code in ["N1V1C1", "N1V4C1BONUS", "N2V52C3", "N2V1C1"]:
                    writer.writerow(["august", code, f"august/{code}"])
            report = audit(path, samples)
            self.assertEqual([c["status"] for c in report["crosswalk"]], ["matched", "unmatched", "unmatched", "ambiguous"])
            self.assertEqual(report["crosswalk"][0]["workbook_rows"], [3])
            self.assertIn("duplicate_fruit_id", report["summary"]["issues_by_code"])

    def test_formula_without_cache_not_evaluated_and_wrong_schema_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.fixture(directory, fruits=[['="F1"', "N1", "V1", "C1", 2, 1, 12, 0]])
            report = audit(path)
            self.assertIsNone(report["fruits"][0]["fruit_id"])
            self.assertEqual(report["fruits"][0]["raw"]["Fruit-ID"], '="F1"')
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "Workbook schema: missing sheet"):
                audit(path)

    def test_output_guard_keeps_source_directories_and_symlink_targets_read_only(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source" / "labels.xlsx"
            samples = root / "collection" / "samples.csv"
            output = root / "output"
            output.mkdir()
            validate_output(output, source, samples)
            for forbidden in [source.parent, source.parent / "nested", samples.parent]:
                with self.assertRaisesRegex(ValueError, "outside source directory"):
                    validate_output(forbidden, source, samples)
            (output / "audit.json").symlink_to(source)
            with self.assertRaisesRegex(ValueError, "must not be a symlink"):
                validate_output(output, source, samples)


if __name__ == "__main__":
    unittest.main()
