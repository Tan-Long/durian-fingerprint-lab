#!/usr/bin/env python3
"""Build the RND to Startup 2026 application for DurianFingerprint."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output/docx/Ban-mo-ta-du-an-DurianFingerprint-RND-to-Startup-2026.docx"
ASSETS = ROOT / "output/docx/assets"

GREEN = "165B45"
TEAL = "158B83"
LIGHT = "EAF3EE"
PALE = "F5F7F4"
GOLD = "D6A640"
GRAY = "5F6F69"
RED = "C94C3B"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def margins(section, top=1.8, bottom=1.6, left=2.0, right=2.0):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def set_cell_text(cell, text, bold=False, color="000000", size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_body(doc, text, bold_lead=None, space_after=5):
    p = doc.add_paragraph(style="Body")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + 0.5 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.add_run("• ")
    p.add_run(text)
    return p


def add_question(doc, text):
    p = doc.add_paragraph(style="Question")
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.add_run(text)
    return p


def add_picture(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=width)
    add_caption(doc, caption)


def add_two_pictures(doc, left_path, left_caption, right_path, right_caption):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, [8.1, 8.1])
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cell, path in zip(table.rows[0].cells, [left_path, right_path]):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(7.7))
    for cell, caption in zip(table.rows[1].cells, [left_caption, right_caption]):
        set_cell_text(cell, caption, bold=True, color=GRAY, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "nil")
                borders.append(el)
            tc_pr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    body = doc.styles.add_style("Body", 1)
    body.base_style = normal
    body.font.name = "Times New Roman"
    body.font.size = Pt(11)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Cm(0.65)
    body.paragraph_format.line_spacing = 1.15

    bullet = doc.styles.add_style("Bullet", 1)
    bullet.base_style = normal
    bullet.font.name = "Times New Roman"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.1

    question = doc.styles.add_style("Question", 1)
    question.base_style = normal
    question.font.name = "Times New Roman"
    question.font.size = Pt(11)
    question.font.bold = True
    question.font.color.rgb = RGBColor.from_string(GREEN)
    question.paragraph_format.space_before = Pt(7)
    question.paragraph_format.space_after = Pt(4)
    question.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.base_style = normal
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)

    for name, size, color in [("Title", 25, GREEN), ("Heading 1", 16, GREEN), ("Heading 2", 13, TEAL)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.text = "DURIANFINGERPRINT | RND TO STARTUP 2026"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.runs[0].font.name = "Times New Roman"
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.add_paragraph())


def add_section_title(doc, number, title):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(f"{number}. {title}")
    return p


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        margins(section)
        add_footer(section)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CUỘC THI KHỞI NGHIỆP SÁNG TẠO\nĐẠI HỌC QUỐC GIA HÀ NỘI\nRND TO STARTUP 2026")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("DURIANFINGERPRINT")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Định danh và truy xuất nguồn gốc sầu riêng đến từng quả")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.paragraph_format.space_after = Pt(12)

    add_picture(doc, ASSETS / "field-capture.jpg", Cm(11.8), "Thu nhận hình ảnh sầu riêng tại vườn bằng thiết bị di động")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Lĩnh vực: Nông nghiệp số - Trí tuệ nhân tạo - Truy xuất nguồn gốc")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NHÓM THỰC HIỆN: SOILTECH")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.add_run("Hà Nội, tháng 9 năm 2026")
    doc.add_page_break()

    # Team page
    doc.add_heading("1. THÔNG TIN NHÓM THỰC HIỆN", level=1)
    table = doc.add_table(rows=3, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["STT", "Họ và tên", "Năm sinh", "Đơn vị", "Chuyên ngành", "Vai trò", "Ghi chú"]
    widths = [0.7, 3.5, 1.4, 2.0, 2.7, 2.3, 1.8]
    set_table_widths(table, widths)
    for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        shade(cell, GREEN)
        set_cell_text(cell, text, True, "FFFFFF", 8.5, WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    rows = [
        ["1", "Nguyễn Ngọc Huyền Quyên", "2006", "SoilTECH", "Khoa học Môi trường", "Nghiên cứu, hiện trường", "Nhóm trưởng"],
        ["2", "Nguyễn Hữu Tân Long", "2004", "SoilTECH", "Khoa học Môi trường", "AI Engineer", ""],
    ]
    for row, values in zip(table.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    info = doc.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [("Số điện thoại liên hệ", "[BỔ SUNG TRƯỚC KHI NỘP]"), ("Email liên hệ", "[BỔ SUNG TRƯỚC KHI NỘP]")]
    for row, (label, value) in zip(info.rows, fields):
        shade(row.cells[0], LIGHT)
        set_cell_text(row.cells[0], label, True, GREEN, 10)
        shade(row.cells[1], "FFF4CC")
        set_cell_text(row.cells[1], value, True, RED, 10)

    doc.add_paragraph()
    add_question(doc, "Năng lực cốt lõi của nhóm")
    add_bullet(doc, "Kết hợp chuyên môn khoa học môi trường với phát triển thuật toán AI và thị giác máy tính.")
    add_bullet(doc, "Đã xây dựng bộ gá quét, quy trình thu nhận nhiều góc, pipeline tách quả, dựng bản đồ bề mặt và tìm vùng tương ứng.")
    add_bullet(doc, "Có khả năng kết nối dữ liệu canh tác, đất và chất lượng sau thu hoạch trong hệ sinh thái SoilTECH.")
    add_picture(doc, ASSETS / "scan-rig.jpg", Cm(9.2), "Bộ gá quét prototype sử dụng điện thoại và hệ thống chiếu sáng")
    doc.add_page_break()

    # Executive summary
    doc.add_heading("2. TÓM TẮT DỰ ÁN", level=1)
    summary = (
        "Sầu riêng là mặt hàng giá trị cao của ngành rau quả Việt Nam, nhưng thông tin truy xuất hiện chủ yếu dừng ở mã vùng trồng, cơ sở đóng gói, lô hàng hoặc tem nhãn. Các lớp thông tin này cần thiết cho quản lý chuỗi cung ứng, song chưa tạo ra dấu hiệu có thể đối chiếu trực tiếp với từng quả trong quá trình thu mua, vận chuyển và bán lẻ. DurianFingerprint đề xuất biến bề mặt tự nhiên của quả thành định danh số. Người trồng hoặc điểm thu mua dùng điện thoại ghi hình quả từ nhiều góc; hệ thống thị giác máy tính phân tích tổ hợp gai, khe, vân vỏ và hình dáng để tạo hồ sơ đặc trưng gắn với vùng trồng, lô hàng và thời điểm thu hoạch. Khi kiểm tra, người dùng chụp lại một vùng vỏ để tìm hồ sơ phù hợp, kể cả khi tem hoặc QR bị mất."
        "\n\nGiải pháp hướng tới nhà vườn, hợp tác xã, cơ sở thu mua - đóng gói, doanh nghiệp xuất khẩu, phòng thử nghiệm và người mua. Prototype gồm bộ gá quét, dữ liệu ảnh thật, bản đồ bề mặt 360 độ phần thân và kết quả đánh dấu đặc trưng trên vỏ. Giai đoạn tiếp theo tập trung kiểm chứng bằng Living Lab, chuẩn hóa quy trình và xây dựng dữ liệu nhiều quả, nhiều điều kiện chụp. Về lâu dài, hồ sơ từng quả có thể liên kết với nhật ký tưới tiêu, phân bón, phân tích đất và chất lượng sau thu hoạch, giúp giảm nhầm lẫn, tăng minh bạch giao dịch và hỗ trợ nhà vườn cải thiện mùa vụ."
    )
    for part in summary.split("\n\n"):
        add_body(doc, part)

    callout = doc.add_table(rows=1, cols=3)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, big, small, color in zip(
        callout.rows[0].cells,
        ["01 QUẢ", "01 DẤU VÂN SỐ", "01 HỒ SƠ"],
        ["Định danh cấp quả", "Từ đặc trưng tự nhiên", "Liên kết chuỗi dữ liệu"],
        [GREEN, TEAL, GOLD],
    ):
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(big)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor.from_string(color)
        p.add_run("\n" + small).font.size = Pt(9)
    doc.add_page_break()

    # Main content
    doc.add_heading("NỘI DUNG CHÍNH CỦA DỰ ÁN", level=0)
    add_section_title(doc, "I", "TÓM TẮT Ý TƯỞNG")
    add_body(doc, "DurianFingerprint phát triển một lớp định danh số cho sầu riêng ở cấp từng quả. Thay vì xem tem hoặc mã QR là bằng chứng duy nhất, giải pháp khai thác cấu trúc tự nhiên của lớp vỏ - tổ hợp gai, khe, vân và hình dáng - như một dấu vân có thể kiểm tra lại bằng hình ảnh. Tại vườn hoặc điểm thu mua, điện thoại ghi hình quả từ nhiều góc. Pipeline thị giác máy tính tách quả khỏi nền, chuẩn hóa các khung hình, xây dựng bản đồ bề mặt và mã hóa vùng đặc trưng. Hồ sơ này được liên kết với thông tin vùng trồng, lô hàng và thời điểm thu hoạch. Khi quả đi qua các khâu vận chuyển, đóng gói hoặc bán lẻ, người dùng chụp lại một phần bề mặt để đối chiếu với dữ liệu đã đăng ký.")
    add_body(doc, "Dự án giải quyết khoảng trống giữa truy xuất theo lô và việc xác minh chính quả đang được giao nhận. Cách tiếp cận này hỗ trợ giảm nhầm lẫn, tráo hàng và tranh chấp nguồn gốc, đồng thời tạo dữ liệu khách quan cho nhà vườn và doanh nghiệp. Ở giai đoạn mở rộng, hồ sơ từng quả có thể liên kết với dữ liệu đất, tưới tiêu, phân bón, thuốc bảo vệ thực vật và kết quả chất lượng sau thu hoạch. Nhờ đó, một thao tác thu nhận không chỉ phục vụ truy xuất mà còn hình thành vòng phản hồi giữa điều kiện canh tác và chất lượng đầu ra. Mục tiêu của giai đoạn thi là hoàn thiện prototype, triển khai pilot có đối chứng và xác định mô hình cung cấp dịch vụ phù hợp cho hợp tác xã, cơ sở đóng gói và doanh nghiệp xuất khẩu.")

    add_section_title(doc, "II", "XÁC ĐỊNH VẤN ĐỀ")
    add_question(doc, "1. Vấn đề cụ thể mà dự án hướng tới là gì?")
    add_body(doc, "Trong 9 tháng đầu năm 2025, xuất khẩu sầu riêng của Việt Nam đạt 2,76 tỷ USD; toàn ngành rau quả đạt 7,759 tỷ USD trong 11 tháng năm 2025 [1][2]. Giá trị lớn khiến nguồn gốc, chất lượng và uy tín của từng mắt xích trong chuỗi trở nên đặc biệt quan trọng. Việt Nam đã xây dựng hệ thống mã số vùng trồng và cơ sở đóng gói phục vụ xuất khẩu; đến ngày 15/8/2025, cả nước có 9.207 mã số vùng trồng và 1.735 mã số cơ sở đóng gói cho các loại quả tươi [3]. Tuy nhiên, các mã này chủ yếu quản lý theo vùng, cơ sở hoặc lô hàng. Khi một quả rời khỏi bao gói hoặc mất tem, việc đối chiếu lại đúng quả với hồ sơ ban đầu vẫn phụ thuộc nhiều vào quy trình vận hành và niềm tin giữa các bên.")
    add_body(doc, "Khoảng trống ở “mét cuối” tạo ra ba nỗi đau: nhà vườn khó chứng minh nguồn gốc và chất lượng của sản phẩm cụ thể; điểm thu mua và doanh nghiệp khó kiểm soát nhầm lẫn hoặc đánh tráo trong quá trình gom hàng; người mua cuối khó kiểm tra thông tin khi tem bị hỏng, mất hoặc thay thế. Song song, dữ liệu canh tác, phân tích đất và chất lượng sau thu hoạch thường nằm ở các biểu mẫu rời rạc, chưa liên kết đến đối tượng vật lý là từng quả.")
    add_question(doc, "2. Vấn đề ảnh hưởng đến những ai và ở đâu?")
    add_bullet(doc, "Nhà vườn và hợp tác xã tại các vùng trồng sầu riêng cần công cụ dễ sử dụng để tạo bằng chứng số và bảo vệ uy tín sản phẩm.")
    add_bullet(doc, "Cơ sở thu mua, đóng gói và doanh nghiệp xuất khẩu cần giảm sai lệch giữa quả thực tế với hồ sơ lô hàng.")
    add_bullet(doc, "Người mua, nhà bán lẻ và người tiêu dùng cần một phương thức kiểm tra trực quan, không phụ thuộc hoàn toàn vào nhãn vật lý.")
    add_bullet(doc, "Phòng thử nghiệm, đơn vị nghiên cứu và cơ quan quản lý địa phương có thể sử dụng dữ liệu tổng hợp để quan sát chất lượng theo khu vực và mùa vụ.")
    doc.add_page_break()

    add_section_title(doc, "III", "GIẢI PHÁP ĐỀ XUẤT")
    add_question(doc, "1. Giải pháp cốt lõi là gì?")
    steps = doc.add_table(rows=5, cols=3)
    steps.style = "Table Grid"
    steps.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(steps, [1.2, 2.6, 12.2])
    data = [
        ("01", "Thu nhận", "Quay hoặc chụp nhiều góc bằng điện thoại; LiDAR là dữ liệu hỗ trợ khi thiết bị có sẵn."),
        ("02", "Chuẩn hóa", "Tách quả khỏi nền, kiểm tra chất lượng ảnh, chọn các góc quan sát phù hợp."),
        ("03", "Tạo dấu vân", "Mã hóa tổ hợp gai, khe, vân vỏ và hình dáng thành hồ sơ đặc trưng của quả."),
        ("04", "Liên kết", "Gắn fruit_id với vùng trồng, lô hàng, thời điểm thu hoạch và dữ liệu canh tác liên quan."),
        ("05", "Xác minh", "Chụp lại một vùng vỏ; hệ thống tìm ứng viên và kiểm tra sự phù hợp hình học."),
    ]
    for row, values in zip(steps.rows, data):
        shade(row.cells[0], GREEN)
        set_cell_text(row.cells[0], values[0], True, "FFFFFF", 11, WD_ALIGN_PARAGRAPH.CENTER)
        shade(row.cells[1], LIGHT)
        set_cell_text(row.cells[1], values[1], True, GREEN, 10)
        set_cell_text(row.cells[2], values[2], size=9.5)

    add_question(doc, "2. Giải pháp mới ở điểm nào?")
    add_bullet(doc, "Định danh ở cấp từng quả bằng chính cấu trúc lớp vỏ, trong khi QR và mã lô đóng vai trò liên kết hồ sơ.")
    add_bullet(doc, "Thiết bị đầu vào là điện thoại phổ biến; người dùng có thể xác minh từ một vùng bề mặt thay vì phải dựng lại mô hình 3D hoàn chỉnh.")
    add_bullet(doc, "Một quy trình dữ liệu phục vụ đồng thời truy xuất, đối chiếu giao nhận và xây nền cho phân tích mối liên hệ giữa canh tác - đất - chất lượng.")
    add_bullet(doc, "Có thể tích hợp qua API với hệ thống quản lý vùng trồng, phòng thử nghiệm, nền tảng SoilTECH hoặc phần mềm của doanh nghiệp.")

    add_question(doc, "3. Minh họa giải pháp")
    add_two_pictures(
        doc,
        ASSETS / "body-texture-map.png",
        "Bản đồ bề mặt vỏ ghép từ nhiều góc chụp",
        ASSETS / "surface-feature-overlay.jpg",
        "Các đặc trưng bề mặt được phát hiện trên ảnh thử nghiệm",
    )
    add_body(doc, "Kết quả xử lý gần nhất tạo được bản đồ bề mặt kích thước 1.259 × 353 px từ 43/48 khung hình hợp lệ, với vùng thân quả được phủ liên tục. Đây là đầu ra kỹ thuật để tiếp tục xây dựng và đánh giá dấu vân số trên bộ dữ liệu nhiều quả.")
    doc.add_page_break()

    add_section_title(doc, "IV", "Ý TƯỞNG KINH DOANH")
    add_question(doc, "1. Khách hàng mục tiêu / người thụ hưởng")
    business = doc.add_table(rows=5, cols=3)
    business.style = "Table Grid"
    business.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(business, [3.2, 5.8, 7.0])
    for cell, text in zip(business.rows[0].cells, ["Nhóm", "Nhu cầu", "Giá trị DurianFingerprint"]):
        shade(cell, GREEN)
        set_cell_text(cell, text, True, "FFFFFF", 9, WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(business.rows[0])
    business_rows = [
        ("Nhà vườn / hợp tác xã", "Chứng minh nguồn gốc, tạo hồ sơ mùa vụ", "Công cụ thu nhận bằng điện thoại; hồ sơ từng quả"),
        ("Cơ sở thu mua / đóng gói", "Giảm nhầm lẫn, kiểm soát giao nhận", "Xác minh quả với hồ sơ đã đăng ký"),
        ("Doanh nghiệp xuất khẩu", "Minh bạch chuỗi cung ứng, quản lý chất lượng", "Dashboard, API và báo cáo theo lô/vùng"),
        ("Phòng thử nghiệm / địa phương", "Liên kết kết quả đo với đối tượng và khu vực", "Kho dữ liệu có cấu trúc, phân tích tổng hợp"),
    ]
    for i, values in enumerate(business_rows, 1):
        if i % 2 == 0:
            for cell in business.rows[i].cells:
                shade(cell, PALE)
        for cell, value in zip(business.rows[i].cells, values):
            set_cell_text(cell, value, size=9)

    add_question(doc, "2. Giá trị và khả năng tự vận hành tài chính")
    add_body(doc, "Mô hình dự kiến là B2B/B2B2C: cung cấp gói khởi tạo dữ liệu và đào tạo vận hành cho hợp tác xã hoặc cơ sở đóng gói; thu phí theo số lượng quả/lô được đăng ký hoặc theo thuê bao phần mềm; cung cấp API và báo cáo dữ liệu cho doanh nghiệp, phòng thử nghiệm và đơn vị quản lý. Trong giai đoạn đầu, dịch vụ triển khai và kiểm chứng hiện trường tạo doanh thu; khi dữ liệu và quy trình ổn định, nền tảng phần mềm giúp giảm chi phí biên và mở rộng sang nhiều vùng trồng.")
    add_question(doc, "3. Kênh tiếp cận khách hàng ban đầu")
    add_bullet(doc, "Pilot cùng một hợp tác xã/cụm nhà vườn và một điểm thu mua - đóng gói tại địa phương trồng sầu riêng.")
    add_bullet(doc, "Mạng lưới SoilTECH, phòng thử nghiệm đất, trường đại học và đối tác nghiên cứu nông nghiệp.")
    add_bullet(doc, "Hội thảo ngành hàng, chương trình đổi mới sáng tạo và hợp tác trực tiếp với doanh nghiệp xuất khẩu.")

    add_section_title(doc, "V", "TÍNH KHẢ THI BAN ĐẦU")
    add_question(doc, "1. Năng lực và tài sản sẵn có")
    add_bullet(doc, "Đội ngũ có nền tảng khoa học môi trường và năng lực phát triển AI/thị giác máy tính.")
    add_bullet(doc, "Bộ gá quét, quy trình thu nhận, dữ liệu ảnh thật và pipeline dựng bản đồ bề mặt đã được hình thành.")
    add_bullet(doc, "Prototype đã phát hiện đặc trưng bề mặt và tìm vùng tương ứng giữa ảnh truy vấn với kho ảnh nhiều góc.")
    add_two_pictures(
        doc,
        ASSETS / "field-capture.jpg",
        "Thử nghiệm thu nhận tại vườn",
        ASSETS / "scan-rig.jpg",
        "Bộ gá thu nhận trong điều kiện kiểm soát",
    )
    add_question(doc, "2. Nguồn lực cần cho giai đoạn tiếp theo")
    add_bullet(doc, "Đối tác hiện trường: nhà vườn/hợp tác xã, điểm thu mua và người dùng kiểm tra độc lập.")
    add_bullet(doc, "Dữ liệu chuẩn hóa: nhiều quả, nhiều buổi chụp, thiết bị và điều kiện ánh sáng; ground truth sau bổ cho bài toán chất lượng.")
    add_bullet(doc, "Kỹ thuật: hoàn thiện ứng dụng thu nhận, dịch vụ đối sánh, lưu trữ an toàn và dashboard quản trị.")
    add_bullet(doc, "Tài chính: thiết bị chiếu sáng - hiệu chuẩn, tổ chức pilot, thu thập/gán nhãn dữ liệu và đánh giá người dùng.")
    add_question(doc, "3. Kế hoạch thử nghiệm và mở rộng")
    add_body(doc, "Pilot đề xuất kéo dài 6 tháng với tối thiểu 30 quả cho bài toán định danh, 60 quả có dữ liệu chất lượng sau bổ và khoảng 20 người dùng thuộc các vai trò nhà vườn, thu mua và vận hành. Thuật toán được khóa trước đợt đánh giá cuối; kết quả được báo cáo theo điều kiện chụp, thiết bị và nhóm người dùng. Sau pilot, giải pháp mở rộng theo ba lớp: tăng số vùng/lô; tích hợp với hệ thống doanh nghiệp qua API; bổ sung mô-đun phân tích đất, canh tác và chất lượng.")
    doc.add_page_break()

    add_section_title(doc, "VI", "TÁC ĐỘNG DỰ KIẾN ĐẾN MÔI TRƯỜNG - XÃ HỘI")
    impact = doc.add_table(rows=5, cols=3)
    impact.style = "Table Grid"
    impact.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(impact, [2.7, 7.2, 6.1])
    for cell, text in zip(impact.rows[0].cells, ["Chiều tác động", "Thay đổi kỳ vọng", "Chỉ số pilot đề xuất"]):
        shade(cell, GREEN)
        set_cell_text(cell, text, True, "FFFFFF", 9, WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(impact.rows[0])
    impact_rows = [
        ("Kinh tế", "Giảm nhầm lẫn/tranh chấp; tăng căn cứ thương lượng cho nhà vườn", "≥70% người dùng đánh giá hồ sơ số hữu ích"),
        ("Xã hội", "Mở rộng khả năng tham gia chuỗi số bằng điện thoại phổ biến", "≥80% người dùng hoàn thành thao tác không cần hỗ trợ trực tiếp"),
        ("Môi trường", "Liên kết đầu vào canh tác với dữ liệu đất và chất lượng để tối ưu theo mùa vụ", "100% mẫu pilot có trường dữ liệu canh tác cốt lõi"),
        ("Kỹ thuật", "Tạo bằng chứng có thể kiểm tra ở cấp từng quả", "Top-1 ≥90%; TAR ≥85% tại FAR ≤1% (mục tiêu pilot)"),
    ]
    for i, values in enumerate(impact_rows, 1):
        if i % 2 == 0:
            for cell in impact.rows[i].cells:
                shade(cell, PALE)
        for cell, value in zip(impact.rows[i].cells, values):
            set_cell_text(cell, value, size=9)
    add_body(doc, "Tác động môi trường của dự án không đến từ việc thay thế ngay một lượng vật tư cố định, mà từ khả năng tạo dữ liệu đủ chi tiết để nhà vườn, phòng thử nghiệm và chuyên gia so sánh điều kiện đất - đầu vào - chất lượng theo mùa vụ. Khi được triển khai ở quy mô lớn, dữ liệu này có thể hỗ trợ điều chỉnh tưới tiêu và phân bón theo nhu cầu thực tế, đồng thời nâng cao trách nhiệm giải trình trong chuỗi cung ứng.")

    add_section_title(doc, "VII", "TÀI LIỆU BỔ SUNG VÀ MINH CHỨNG")
    evidence = doc.add_table(rows=5, cols=3)
    evidence.style = "Table Grid"
    evidence.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(evidence, [3.0, 6.6, 6.4])
    for cell, text in zip(evidence.rows[0].cells, ["Minh chứng", "Nội dung thể hiện", "Vai trò trong dự án"]):
        shade(cell, GREEN)
        set_cell_text(cell, text, True, "FFFFFF", 9, WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(evidence.rows[0])
    evidence_rows = [
        ("Video thu nhận tại vườn", "Người dùng chụp quả bằng điện thoại trong bối cảnh thật", "Kiểm tra khả năng triển khai hiện trường"),
        ("Bộ gá quét", "Điện thoại, chiếu sáng và quả được bố trí để thu nhiều góc", "Tạo dữ liệu kiểm soát phục vụ phát triển thuật toán"),
        ("Bản đồ bề mặt", "Ảnh lớp vỏ được ghép thành dải quan sát 360 độ phần thân", "Nền tảng xây dựng dấu vân số"),
        ("Ảnh đặc trưng", "Các điểm/gai nổi bật được phát hiện trên bề mặt", "Đầu vào cho đối sánh và kiểm tra hình học"),
    ]
    for i, values in enumerate(evidence_rows, 1):
        if i % 2 == 0:
            for cell in evidence.rows[i].cells:
                shade(cell, PALE)
        for cell, value in zip(evidence.rows[i].cells, values):
            set_cell_text(cell, value, size=9)

    add_picture(doc, ASSETS / "body-texture-map.png", Cm(15.5), "Minh chứng prototype: bản đồ bề mặt lớp vỏ từ dữ liệu thu nhận nhiều góc")

    doc.add_heading("Nguồn số liệu tham khảo", level=2)
    sources = [
        ("[1] Bộ Công Thương, Bản tin thị trường nông, lâm, thủy sản quý III/2025", "https://moit.gov.vn/upload/2005517/fck/files/B___n_tin_Th____tr_____ng_NLTS_T___ng_k___t_qu___III_ra_ng__y_31_10_2025__1__5d2a6.pdf"),
        ("[2] Bộ Công Thương, Bản tin thị trường nông, lâm, thủy sản ngày 22/12/2025", "https://moit.gov.vn/upload/2005517/fck/files/B___n_tin_Th____tr_____ng_NLTS_s____ra_ng__y_22_12_2025__1__66959.pdf"),
        ("[3] Cục Trồng trọt và Bảo vệ thực vật, số liệu mã số vùng trồng và cơ sở đóng gói đến 15/8/2025", "https://ppd.gov.vn/tin-tuc-noi-bat/hoi-nghi-so-ket-san-xuat-trong-trot-nam-2025-trien-khai-ke-hoach--vu-dong-2025-va-cac-vu-san-xuat-nam-2026-tai-cac-tinh-phia-bac.html"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="Body")
        p.paragraph_format.first_line_indent = Cm(0)
        add_hyperlink(p, label, url)

    doc.add_paragraph()
    note = doc.add_table(rows=1, cols=1)
    shade(note.cell(0, 0), LIGHT)
    set_cell_text(note.cell(0, 0), "Tư liệu hình ảnh và video prototype gốc do nhóm DurianFingerprint/SoilTECH cung cấp, tháng 9/2026.", True, GREEN, 9.5, WD_ALIGN_PARAGRAPH.CENTER)

    # Keep document metadata and save.
    doc.core_properties.title = "Bản mô tả dự án DurianFingerprint - RND to Startup 2026"
    doc.core_properties.subject = "Định danh và truy xuất nguồn gốc sầu riêng đến từng quả"
    doc.core_properties.author = "Nhóm DurianFingerprint / SoilTECH"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
