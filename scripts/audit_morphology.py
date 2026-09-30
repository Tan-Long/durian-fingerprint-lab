#!/usr/bin/env python3
"""Audit morphology labels without editing sources (requires openpyxl).

Example: python3 scripts/audit_morphology.py /path/to/source/SAMPLE_Fruit_morphology.xlsx \
    --samples /path/to/manifests/samples.csv --output output/morphology

Writes audit.json (original/cached cells, per-target checks, issues and crosswalk)
and sample_crosswalk.csv. 'valid' means passes these label checks, not reviewed
ground truth or training-ready. Formula values use Excel's saved cache; this
tool neither evaluates formulas nor repairs labels. Missing values remain null.
Crosswalk 'status' is the mapping result; 'sample_status' preserves CSV status.
Output must be outside source directories and may not overwrite symlink targets.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FRUIT_SHEET = "Fruit-list"
LOCULE_SHEET = "Thông tin từng hộc"
TARGETS = {
    "locule_count": ("Số lượng hộc", True, False),
    "aril_count": ("Tổng số múi", True, True),
    "shell_thickness_mm": ("dày vỏ", False, False),
    "empty_locule_count": ("Số hộc lép", True, True),
}
REQUIRED = {
    FRUIT_SHEET: ["Fruit-ID", "Group-ID", "Farm-ID", "Tree-ID", *[t[0] for t in TARGETS.values()]],
    LOCULE_SHEET: ["Fruit-ID", "LoculeNo (hộc)", "ArilNumber (số múi)"],
}


def numeric_problem(value, integer=False, allow_zero=False):
    if value is None or value == "":
        return "missing_numeric"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return "nonnumeric"
    if value < 0 or (value == 0 and not allow_zero):
        return "negative_numeric" if allow_zero else "nonpositive_numeric"
    if integer and int(value) != value:
        return "noninteger_numeric"
    return None


def audit(workbook_path: Path, samples_path: Path | None = None) -> dict:
    issues = []

    def issue(code, records, fields, **detail):
        issues.append({"code": code, "sources": [
            {"sheet": record["sheet"], "row": record["row"],
             "cells": [record["cells"][field] for field in fields]}
            for record in records
        ], **detail})

    raw_book = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        cached_book = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            sheets = {}
            for name, required in REQUIRED.items():
                if name not in raw_book.sheetnames:
                    raise ValueError(f"Workbook schema: missing sheet {name!r}")
                raw_sheet = raw_book[name]
                header = next(raw_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                missing = [field for field in required if header.count(field) != 1]
                if missing:
                    raise ValueError(f"Workbook schema: {name!r} requires unique headers {missing}")
                # The supplied template has headers at row 1 and units at row 2.
                # Ignore the separate glossary starting in column V of Fruit-list.
                width = 20 if name == FRUIT_SHEET else len(header)
                columns = {field: i for i, field in enumerate(header[:width]) if field is not None}
                if not set(required) <= columns.keys():
                    raise ValueError(f"Workbook schema: {name!r} label columns outside template data area")
                if name == FRUIT_SHEET:
                    units = next(raw_sheet.iter_rows(min_row=2, max_row=2, max_col=width, values_only=True))
                    if units[columns["dày vỏ"]] != "mm":
                        raise ValueError("Workbook schema: 'dày vỏ' unit in row 2 must be 'mm'; no conversion inferred")
                records = []
                for row_number, (raw, values) in enumerate(zip(
                    raw_sheet.iter_rows(min_row=3, max_col=width, values_only=True),
                    cached_book[name].iter_rows(min_row=3, max_col=width, values_only=True),
                ), 3):
                    if not any(value is not None for value in raw):
                        continue
                    record = {
                        "sheet": name, "row": row_number,
                        "values": {field: values[i] for field, i in columns.items()},
                        "raw": {field: raw[i] for field, i in columns.items()},
                        "cells": {field: f"{get_column_letter(i + 1)}{row_number}" for field, i in columns.items()},
                    }
                    record["fruit_id"] = record["values"]["Fruit-ID"]
                    records.append(record)
                sheets[name] = records
        finally:
            cached_book.close()
    finally:
        raw_book.close()

    fruits, locules = sheets[FRUIT_SHEET], sheets[LOCULE_SHEET]
    by_fruit = defaultdict(list)
    by_detail = defaultdict(list)
    for records, grouped in [(fruits, by_fruit), (locules, by_detail)]:
        for record in records:
            fruit_id = record["fruit_id"]
            if not isinstance(fruit_id, str) or not fruit_id.strip():
                issue("missing_or_invalid_fruit_id", [record], ["Fruit-ID"])
            else:
                grouped[fruit_id].append(record)
    for fruit_id, records in by_fruit.items():
        if len(records) > 1:
            issue("duplicate_fruit_id", records, ["Fruit-ID"], fruit_id=fruit_id)

    for locule in locules:
        locule["checks"] = {}
        for field, allow_zero in [("LoculeNo (hộc)", False), ("ArilNumber (số múi)", True)]:
            problem = numeric_problem(locule["values"][field], integer=True, allow_zero=allow_zero)
            locule["checks"][field] = problem or "valid"
            if problem:
                issue(problem, [locule], [field], field=field)
        if locule["fruit_id"] not in by_fruit:
            issue("unknown_detail_fruit_id", [locule], ["Fruit-ID"])

    duplicated_keys = set()
    for fruit_id, records in by_detail.items():
        by_number = defaultdict(list)
        for record in records:
            if record["checks"]["LoculeNo (hộc)"] == "valid":
                by_number[record["values"]["LoculeNo (hộc)"]].append(record)
        for number, duplicates in by_number.items():
            if len(duplicates) > 1:
                duplicated_keys.add(fruit_id)
                issue("duplicate_locule_key", duplicates, ["Fruit-ID", "LoculeNo (hộc)"],
                      fruit_id=fruit_id, locule_number=number)

    for fruit in fruits:
        fruit_id = fruit["fruit_id"]
        fruit["targets"] = {}
        for target, (field, integer, allow_zero) in TARGETS.items():
            value = fruit["values"][field]
            problem = numeric_problem(value, integer, allow_zero)
            reasons = [problem] if problem else []
            if len(by_fruit.get(fruit_id, [])) != 1:
                reasons.append("unresolved_fruit_identity")
            fruit["targets"][target] = {"value": value, "reasons": reasons}
            if problem:
                issue(problem, [fruit], [field], target=target)

        details = by_detail.get(fruit_id, [])
        numbers = [d["values"]["LoculeNo (hộc)"] for d in details
                   if d["checks"]["LoculeNo (hộc)"] == "valid"]
        keys_ok = bool(details) and len(numbers) == len(details) and fruit_id not in duplicated_keys
        fruit["detail_summary"] = {"rows": [d["row"] for d in details],
                                   "row_count": len(details), "unique_locule_count": len(set(numbers)),
                                   "locule_numbers": numbers, "aril_sum": None}
        if not keys_ok:
            issue("unresolved_locule_keys", [fruit], ["Fruit-ID"], detail_rows=fruit["detail_summary"]["rows"])
            for target in ["locule_count", "aril_count"]:
                fruit["targets"][target]["reasons"].append("unresolved_locule_keys")
        expected_count = fruit["targets"]["locule_count"]
        if not numeric_problem(expected_count["value"], True):
            if len(set(numbers)) != expected_count["value"]:
                issue("locule_count_mismatch", [fruit], ["Số lượng hộc"],
                      expected=expected_count["value"], observed_unique=len(set(numbers)), observed_rows=len(details))
                expected_count["reasons"].append("locule_count_mismatch")
                fruit["targets"]["aril_count"]["reasons"].append("incomplete_or_conflicting_locules")
            if keys_ok and (len(numbers) != expected_count["value"] or
                            any(number != index for index, number in enumerate(sorted(numbers), 1))):
                issue("locule_number_sequence_mismatch", details, ["LoculeNo (hộc)"])
                expected_count["reasons"].append("locule_number_sequence_mismatch")
                fruit["targets"]["aril_count"]["reasons"].append("locule_number_sequence_mismatch")
        arils_ok = bool(details) and all(d["checks"]["ArilNumber (số múi)"] == "valid" for d in details)
        if arils_ok:
            total = sum(d["values"]["ArilNumber (số múi)"] for d in details)
            fruit["detail_summary"]["aril_sum"] = total
            target = fruit["targets"]["aril_count"]
            if not numeric_problem(target["value"], True, True) and total != target["value"]:
                issue("aril_sum_mismatch", [fruit], ["Tổng số múi"], expected=target["value"], observed=total)
                target["reasons"].append("aril_sum_mismatch")
        else:
            fruit["targets"]["aril_count"]["reasons"].append("missing_or_invalid_detail_arils")
        for target in fruit["targets"].values():
            target["status"] = "review" if target["reasons"] else "valid"

    crosswalk = []
    if samples_path is not None:
        with samples_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {"collection_id", "sample_id", "sample_uid"}
            if not required <= set(reader.fieldnames or []):
                raise ValueError(f"Samples schema: required columns {sorted(required)}")
            samples = list(reader)
        candidates = defaultdict(list)
        for fruit in fruits:
            parts = [fruit["values"][field] for field in ["Group-ID", "Farm-ID", "Tree-ID"]]
            if all(isinstance(part, str) and re.fullmatch(pattern, part)
                   for part, pattern in zip(parts, [r"N\d+", r"V\d+", r"C\d+"])):
                candidates["".join(parts)].append(fruit)
        sample_counts = Counter(sample["sample_id"] for sample in samples)
        for csv_row, sample in enumerate(samples, 2):
            matches = candidates.get(sample["sample_id"], [])
            status = "unmatched"
            if sample_counts[sample["sample_id"]] > 1 or len(matches) > 1:
                status = "ambiguous"
            elif len(matches) == 1:
                status = "matched" if len(by_fruit.get(matches[0]["fruit_id"], [])) == 1 else "ambiguous"
            crosswalk.append({**sample, "samples_csv_row": csv_row,
                              "sample_status": sample.get("status"), "status": status,
                              "fruit_id": matches[0]["fruit_id"] if status == "matched" else None,
                              "workbook_sheet": FRUIT_SHEET,
                              "workbook_rows": [fruit["row"] for fruit in matches],
                              "mapping_cells": [{field: fruit["cells"][field] for field in
                                                 ["Group-ID", "Farm-ID", "Tree-ID", "Fruit-ID"]}
                                                for fruit in matches]})

    return {
        "schema_version": 1,
        "sources": {"workbook": str(workbook_path.resolve()),
                    "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                    "samples": str(samples_path.resolve()) if samples_path else None,
                    "samples_sha256": hashlib.sha256(samples_path.read_bytes()).hexdigest() if samples_path else None},
        "limits": "Read-only label checks, not reviewed ground truth or training readiness. Formula caches may be stale. No ID repair or model evaluation.",
        "summary": {"fruit_rows": len(fruits), "locule_rows": len(locules),
                    "locule_rows_with_fruit_id": sum(len(rows) for rows in by_detail.values()),
                    "locule_rows_without_valid_fruit_id": len(locules) - sum(len(rows) for rows in by_detail.values()),
                    "valid_locule_number_rows": sum(l["checks"]["LoculeNo (hộc)"] == "valid" for l in locules),
                    "valid_locule_number_rows_with_fruit_id": sum(
                        l["checks"]["LoculeNo (hộc)"] == "valid" for rows in by_detail.values() for l in rows),
                    "valid_targets": {t: sum(f["targets"][t]["status"] == "valid" for f in fruits) for t in TARGETS},
                    "issues_by_code": dict(Counter(i["code"] for i in issues)),
                    "crosswalk_status": dict(Counter(c["status"] for c in crosswalk))},
        "fruits": fruits, "locules": locules, "issues": issues, "crosswalk": crosswalk,
    }


def validate_output(output: Path, workbook: Path, samples: Path | None = None):
    for source in [workbook, *([samples] if samples else [])]:
        if output.resolve().is_relative_to(source.resolve().parent):
            raise ValueError(f"Output must be outside source directory: {source.resolve().parent}")
    for name in ["audit.json", "sample_crosswalk.csv"]:
        if (output / name).is_symlink():
            raise ValueError(f"Output file must not be a symlink: {output / name}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--samples", type=Path, help="Optional collection manifests/samples.csv")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        validate_output(args.output, args.workbook, args.samples)
        report = audit(args.workbook, args.samples)
    except (OSError, ValueError) as error:
        parser.error(str(error))
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    if args.samples:
        fields = ["collection_id", "sample_id", "sample_uid", "sample_status", "samples_csv_row", "status", "fruit_id", "workbook_sheet", "workbook_rows", "mapping_cells"]
        with (args.output / "sample_crosswalk.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows({**row, "workbook_rows": json.dumps(row["workbook_rows"]),
                              "mapping_cells": json.dumps(row["mapping_cells"])} for row in report["crosswalk"])
    print(json.dumps(report["summary"], ensure_ascii=False))
    print(f"Audit: {args.output / 'audit.json'}")


if __name__ == "__main__":
    main()
